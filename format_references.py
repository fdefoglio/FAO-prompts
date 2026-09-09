import re
import sys
import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def format_references(doc_path):
    # Create output filename
    base, ext = os.path.splitext(doc_path)
    output_path = f"{base}_formatted{ext}"
    
    # Open the document
    doc = Document(doc_path)
    
    # Find the References section
    references_section = False
    references_start = None
    references_end = None
    
    # First, identify where the References section starts and ends
    for i, para in enumerate(doc.paragraphs):
        if para.text.strip().lower() == "references" or para.text.strip().lower() == "references ==========":
            references_start = i
            references_section = True
        elif references_section and (para.text.strip().lower() == "further reading" or para.text.strip().lower() == "further reading ================"):
            references_end = i
            break
    
    if references_start is None:
        print("References section not found!")
        return
    
    if references_end is None:
        references_end = len(doc.paragraphs)
    
    # Process each reference
    for i in range(references_start + 1, references_end):
        para = doc.paragraphs[i]
        text = para.text.strip()
        
        # Skip empty paragraphs
        if not text:
            continue
            
        # Pattern to match author(s) and date - case 1: both not bolded (including n.d. and years with letters)
        # Updated to be more flexible with spacing and punctuation
        pattern1 = r'^([^:]+(?: & [^:]+)*)\s+((?:\d{4}[a-z]?|n\.d\.))\.'
        
        # Pattern to match author(s) bolded but date not bolded - case 2 (including n.d. and years with letters)
        pattern2 = r'^(\*\*[^:]+(?: & [^:]+)*\*\*)\s+((?:\d{4}[a-z]?|n\.d\.))\.'
        
        # Try pattern2 first (author bolded, date not)
        match = re.match(pattern2, text)
        
        if match:
            # Clear the paragraph
            para.text = ""
            
            # Add the bolded author(s) and date
            run = para.add_run(match.group(1) + " " + match.group(2) + ". ")
            run.bold = True
            
            # Add the rest of the reference
            para.add_run(text[len(match.group(0)):])
        else:
            # Try pattern1 (neither bolded)
            match = re.match(pattern1, text)
            
            if match:
                # Clear the paragraph
                para.text = ""
                
                # Add the bolded author(s) and date
                run = para.add_run(match.group(1) + " " + match.group(2) + ". ")
                run.bold = True
                
                # Add the rest of the reference
                para.add_run(text[len(match.group(0)):])
            else:
                # Try pattern for cases with only authors and no date
                alt_pattern = r'^(\*\*[^:]+(?: & [^:]+)*\*\*)'
                alt_match = re.match(alt_pattern, text)
                
                if alt_match:
                    # Clear the paragraph
                    para.text = ""
                    
                    # Add the bolded author(s)
                    run = para.add_run(alt_match.group(1) + " ")
                    run.bold = True
                    
                    # Add the rest of the reference
                    para.add_run(text[len(alt_match.group(0)):])
    
    # Save the document with a new name
    doc.save(output_path)
    print(f"References formatted successfully in {output_path}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        doc_path = sys.argv[1]
    else:
        doc_path = "REF1786799717_P1_final_Edited_v2.docx"
    
    format_references(doc_path)